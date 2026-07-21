from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.config import RESULTS_DIR
from src.data import load_dataset
from src.thesis_results import build_table_4_1
from src.thesis_text import (
    apply_global_content_fixes,
    build_paragraph_updates,
    build_table_4_1_rows,
    build_table_4_2_rows,
    dehyphenate_text,
    load_results_from_disk,
)

THESIS = Path(r"c:\Users\user\Desktop\code project masters\extras\THESIS_OBANOR_TEHILLA_49_.docx")
BACKUP = THESIS.with_name(
    f"THESIS_OBANOR_TEHILLA_49_BACKUP_{datetime.now():%Y%m%d_%H%M%S}.docx"
)


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def add_table_column(table) -> None:
    for row in table.rows:
        row._tr.append(deepcopy(row.cells[-1]._tc))


def patch_table(table, rows: list[list[str]]) -> None:
    needed_cols = max(len(row) for row in rows)
    while len(table.rows[0].cells) < needed_cols:
        add_table_column(table)

    for row_idx, values in enumerate(rows):
        for col_idx, value in enumerate(values):
            table.rows[row_idx].cells[col_idx].text = value


def patch_thesis_document(
    country_df: pd.DataFrame | None = None,
    temporal_df: pd.DataFrame | None = None,
    summary: dict | None = None,
    table_4_1: pd.DataFrame | None = None,
) -> None:
    if country_df is None or temporal_df is None or summary is None:
        country_df, temporal_df, summary = load_results_from_disk()

    if table_4_1 is None:
        table_4_1 = build_table_4_1()

    df = load_dataset()
    default_rate = df["Default"].mean()
    train_countries = summary["country_split"]["train_countries"]
    test_countries = summary["country_split"]["test_countries"]
    probit_detail = {
        "pseudo_r2": summary["country_split"]["probit_pseudo_r2"],
        "llr_pvalue": summary["country_split"]["probit_llr_pvalue"],
    }

    paragraph_updates = build_paragraph_updates(
        country_df,
        temporal_df,
        probit_detail,
        default_rate,
        len(train_countries),
        len(test_countries),
    )

    shutil.copy2(THESIS, BACKUP)
    doc = Document(THESIS)

    for idx, text in paragraph_updates.items():
        if idx < len(doc.paragraphs):
            set_paragraph_text(doc.paragraphs[idx], text)

    for idx, para in enumerate(doc.paragraphs):
        if idx >= 690:
            continue
        text = para.text.strip()
        if not text:
            continue
        updated = apply_global_content_fixes(text, default_rate)
        if updated != text:
            set_paragraph_text(para, updated)
            text = updated
        if "-" in text or "–" in text or "—" in text:
            set_paragraph_text(para, dehyphenate_text(text))

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if not text:
                    continue
                updated = apply_global_content_fixes(text, default_rate)
                if updated != text:
                    cell.text = updated
                    text = updated
                if "-" in text or "–" in text:
                    cell.text = dehyphenate_text(text)

    patch_table(doc.tables[5], build_table_4_1_rows(table_4_1))
    patch_table(doc.tables[6], build_table_4_2_rows(country_df, probit_detail))

    doc.save(THESIS)
    print(f"Backup saved to: {BACKUP}")
    print(f"Updated thesis: {THESIS}")


def main() -> None:
    patch_thesis_document()


if __name__ == "__main__":
    main()
